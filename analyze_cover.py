# -*- coding: utf-8 -*-
from docx import Document

print("=== 分析封面模板结构 ===")
doc = Document('模板文件/template_报告封面.docx')

print(f"\n段落总数: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if text:
        print(f"段落{i}: \"{text[:40]}...\" (runs={len(p.runs)})")

print(f"\n表格总数: {len(doc.tables)}")
for ti, table in enumerate(doc.tables):
    print(f"\n表格{ti}:")
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            text = cell.text.strip()
            if '标题' in text or '项目名称' in text or len(text) > 20:
                print(f"  行{ri}列{ci}: \"{text[:40]}\"")

try:
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn
    print(f"\n检查文本框/形状...")
    for shape in doc.inline_shapes:
        print(f"  内嵌形状: {shape}")
except:
    print("\n无法检查形状（可能需要额外库）")

